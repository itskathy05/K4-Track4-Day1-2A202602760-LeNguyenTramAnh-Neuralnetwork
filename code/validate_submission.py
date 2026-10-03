"""Strict, read-only validator and packager for the final submission."""
from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


REQUIRED = (
    "REPORT.md", "experiments.xlsx", "predictions_eval.csv", "eval_result.json",
    "figures", "results", "code/lab.ipynb",
)
BANNED_SUFFIXES = {".pt", ".pth", ".ckpt", ".npz", ".gz"}


def validate_submission(root: str | Path, eval_npz: str | Path | None = None) -> list[str]:
    root = Path(root)
    errors: list[str] = []
    for relative in REQUIRED:
        if not (root / relative).exists():
            errors.append(f"Thiếu {relative}")
    if errors:
        return errors

    banned = [p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in BANNED_SUFFIXES]
    banned += [p for p in root.rglob("*") if "__pycache__" in p.parts or ".ipynb_checkpoints" in p.parts]
    if banned:
        errors.append("Có artifact bị cấm: " + ", ".join(str(p.relative_to(root)) for p in banned[:10]))

    for py in (root / "code").rglob("*.py"):
        text = py.read_text(encoding="utf-8")
        unfinished_raise = "raise " + "NotImplemented" + "Error"
        unfinished_comment = r"#\s*" + "TO" + "DO"
        if unfinished_raise in text or re.search(unfinished_comment, text, re.I):
            errors.append(f"Code chưa hoàn thiện: {py.relative_to(root)}")

    report = (root / "REPORT.md").read_text(encoding="utf-8")
    if re.search(r"___|<MSSV>|<Họ tên>|TODO", report, re.I):
        errors.append("REPORT.md còn placeholder")

    wb = openpyxl.load_workbook(root / "experiments.xlsx", data_only=False)
    if wb.sheetnames != ["Legend", "Experiments", "Seeds", "Summary"]:
        errors.append(f"Sai sheet Excel: {wb.sheetnames}")
    ws = wb["Experiments"]
    headers = [c.value for c in ws[1]]
    exp_col = headers.index("exp_id") + 1
    fig_col = headers.index("figure_file") + 1
    experiment_rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, exp_col).value]
    exp_ids = [str(ws.cell(r, exp_col).value) for r in experiment_rows]
    if len(exp_ids) != len(set(exp_ids)):
        errors.append("exp_id trong Excel bị trùng")
    for row in experiment_rows:
        exp_id = str(ws.cell(row, exp_col).value)
        figure_file = ws.cell(row, fig_col).value
        if figure_file != f"figures/{exp_id}.png" or not (root / str(figure_file)).exists():
            errors.append(f"Thiếu/sai ảnh của {exp_id}")
    result_ids = {p.stem for p in (root / "results").glob("*.json")}
    if set(exp_ids) != result_ids:
        errors.append("Danh sách exp_id trong Excel không khớp results/*.json")

    pred = pd.read_csv(root / "predictions_eval.csv")
    if list(pred.columns) != ["row_id", "pred"]:
        errors.append("predictions_eval.csv phải có đúng cột row_id,pred")
    else:
        if len(pred) != 116_203 or pred["row_id"].duplicated().any():
            errors.append("predictions_eval.csv không đủ 116.203 row_id duy nhất")
        if pred.isna().any().any() or not pred["pred"].isin(range(7)).all():
            errors.append("predictions_eval.csv có giá trị trống hoặc nhãn ngoài 0..6")

    eval_result = json.loads((root / "eval_result.json").read_text(encoding="utf-8"))
    if eval_result.get("n_eval") != 116_203 or not {"accuracy", "macro_f1", "per_class", "confusion_matrix"}.issubset(eval_result):
        errors.append("eval_result.json thiếu trường hoặc sai n_eval")
    if eval_npz is not None and Path(eval_npz).exists() and not errors:
        with np.load(eval_npz, allow_pickle=False) as ev:
            truth = pd.Series(ev["y"], index=ev["row_id"])
        aligned = pred.set_index("row_id")["pred"].loc[truth.index].to_numpy()
        y = truth.to_numpy()
        cm = np.zeros((7, 7), dtype=np.int64)
        np.add.at(cm, (y, aligned), 1)
        tp = np.diag(cm).astype(float)
        precision = np.divide(tp, cm.sum(0), out=np.zeros(7), where=cm.sum(0) > 0)
        recall = np.divide(tp, cm.sum(1), out=np.zeros(7), where=cm.sum(1) > 0)
        f1 = np.divide(2 * precision * recall, precision + recall, out=np.zeros(7), where=precision + recall > 0)
        acc, macro = float(tp.sum() / cm.sum()), float(f1.mean())
        if abs(acc - eval_result["accuracy"]) > 5e-4 or abs(macro - eval_result["macro_f1"]) > 5e-4:
            errors.append("Metric recompute không khớp eval_result.json")
    return errors


def create_zip(root: str | Path, out_path: str | Path) -> Path:
    root, out_path = Path(root), Path(out_path)
    errors = validate_submission(root)
    if errors:
        raise ValueError("Submission chưa hợp lệ:\n- " + "\n- ".join(errors))
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for file in root.rglob("*"):
            if file.is_file() and file.suffix.lower() not in BANNED_SUFFIXES and "__pycache__" not in file.parts:
                archive.write(file, Path(root.name) / file.relative_to(root))
    return out_path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("root")
    parser.add_argument("--eval-npz", default=None)
    parser.add_argument("--zip", dest="zip_path", default=None)
    args = parser.parse_args()
    errors = validate_submission(args.root, args.eval_npz)
    if errors:
        print("SUBMISSION CHƯA HỢP LỆ:\n- " + "\n- ".join(errors))
        raise SystemExit(1)
    print("SUBMISSION HỢP LỆ")
    if args.zip_path:
        print("Đã tạo", create_zip(args.root, args.zip_path))


if __name__ == "__main__":
    main()
